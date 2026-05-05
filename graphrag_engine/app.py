import io
import re

import streamlit as st
import tiktoken
from pypdf import PdfReader
from docx import Document
from openai import OpenAI
from rdflib import Graph, RDF, OWL
from poml import poml
from loguru import logger

from config import (
    OPENAI_API_KEY,
    KNOWN_MODEL_LIMITS,
    TIKTOKEN_ENCODING,
    BATCH_TOKEN_BUDGET,
    CHUNK_TOKEN_LIMIT,
    ACCEPTED_FILE_TYPES,
    PROMPTS_DIR,
)
from graph import (
    Neo4jClient,
    ontology_to_schema,
    build_knowledge_graph,
    enrich_relationships,
    normalize_entity_names,
    query_graph_rag,
    get_graph_stats,
    find_duplicate_entities,
    has_any_entities,
    compute_edge_weights,
    create_evidence_layer,
    create_web_evidence,
    aggregate_evidence,
    enrich_temporal,
    get_evidence_stats,
    clear_evidence_layer,
)
from entity_resolution import (
    find_candidate_pairs,
    score_exact,
    score_embedding_batch,
    score_llm_batch,
    build_transitive_clusters,
    merge_cluster,
    undo_merge,
    find_orphaned_nodes,
)
from web_sources import (
    extract_topics,
    search_and_fetch,
    build_web_knowledge_graph,
    compute_similar_to_edges,
    remove_web_content,
    get_web_source_stats,
)

ENCODING = tiktoken.get_encoding(TIKTOKEN_ENCODING)


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()

    if name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8")

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(uploaded_file.read()))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(uploaded_file.read()))
        return "\n".join(para.text for para in doc.paragraphs)

    return ""


# ---------------------------------------------------------------------------
# Token helpers
# ---------------------------------------------------------------------------
def count_tokens(text: str) -> int:
    return len(ENCODING.encode(text))


# ---------------------------------------------------------------------------
# Chunking — general-purpose paragraph splitting
# ---------------------------------------------------------------------------
def chunk_text(text: str) -> list[str]:
    paragraphs = re.split(r"\n\n+", text.strip())
    paragraphs = [p.strip() for p in paragraphs if p.strip()]

    chunks: list[str] = []
    current_parts: list[str] = []
    current_tokens = 0

    for para in paragraphs:
        para_tokens = count_tokens(para)

        if para_tokens > CHUNK_TOKEN_LIMIT:
            if current_parts:
                chunks.append("\n\n".join(current_parts))
                current_parts = []
                current_tokens = 0
            chunks.append(para)
            logger.warning(f"Single paragraph exceeds chunk limit ({para_tokens} tokens)")
            continue

        if current_tokens + para_tokens > CHUNK_TOKEN_LIMIT and current_parts:
            chunks.append("\n\n".join(current_parts))
            current_parts = []
            current_tokens = 0

        current_parts.append(para)
        current_tokens += para_tokens

    if current_parts:
        chunks.append("\n\n".join(current_parts))

    return chunks


def chunk_all_documents(docs: list[dict]) -> list[str]:
    all_chunks: list[str] = []
    for doc in docs:
        file_chunks = chunk_text(doc["text"])
        all_chunks.extend(file_chunks)
    return all_chunks


# ---------------------------------------------------------------------------
# Batching — greedy bin-packing
# ---------------------------------------------------------------------------
def batch_chunks(chunks: list[str]) -> list[list[str]]:
    batches: list[list[str]] = []
    current_batch: list[str] = []
    current_tokens = 0

    for chunk in chunks:
        chunk_tokens = count_tokens(chunk)
        if current_tokens + chunk_tokens > BATCH_TOKEN_BUDGET and current_batch:
            batches.append(current_batch)
            current_batch = []
            current_tokens = 0
        current_batch.append(chunk)
        current_tokens += chunk_tokens

    if current_batch:
        batches.append(current_batch)

    return batches


# ---------------------------------------------------------------------------
# TTL helpers
# ---------------------------------------------------------------------------
def strip_code_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:turtle|ttl|sparql)?\s*\n", "", text)
    text = re.sub(r"\n```\s*$", "", text)
    return text.strip()


def validate_ttl(ttl_string: str) -> tuple[bool, str]:
    try:
        g = Graph()
        g.parse(data=ttl_string, format="turtle")
        return True, g.serialize(format="turtle")
    except Exception as e:
        return False, str(e)


def merge_ttl_fragments(fragments: list[str]) -> str:
    merged = Graph()
    for frag in fragments:
        try:
            merged.parse(data=frag, format="turtle")
        except Exception as e:
            logger.warning(f"Skipping invalid fragment during merge: {e}")
    return merged.serialize(format="turtle")


def extract_registry(ttl_string: str) -> set[str]:
    g = Graph()
    g.parse(data=ttl_string, format="turtle")
    names: set[str] = set()
    for rdf_type in (OWL.Class, OWL.ObjectProperty, OWL.DatatypeProperty):
        for s in g.subjects(RDF.type, rdf_type):
            local_name = str(s).split("#")[-1].split("/")[-1]
            if local_name:
                names.add(local_name)
    return names


# ---------------------------------------------------------------------------
# LLM calls
# ---------------------------------------------------------------------------
def call_prompt_1a(client: OpenAI, model: str, doc_text: str, registry_text: str = "") -> str:
    params = poml(
        str(PROMPTS_DIR / "ontology_prompt.poml"),
        context={"documents": doc_text, "registry": registry_text},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    raw = response.choices[0].message.content
    return strip_code_fences(raw)


def call_prompt_1b(client: OpenAI, model: str, doc_text: str, existing_ttl: str) -> str:
    params = poml(
        str(PROMPTS_DIR / "refine_ontology.poml"),
        context={"documents": doc_text, "existing_ttl": existing_ttl},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    return strip_code_fences(response.choices[0].message.content)


def call_prompt_2(client: OpenAI, model: str, ttl_content: str, errors: str = "") -> str:
    params = poml(
        str(PROMPTS_DIR / "validate_syntax.poml"),
        context={"ttl_content": ttl_content, "errors": errors},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    return strip_code_fences(response.choices[0].message.content)


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------
def build_document_block(docs: list[dict]) -> str:
    parts = []
    for i, doc in enumerate(docs, 1):
        parts.append(f"--- Document {i}: {doc['name']} ---\n{doc['text']}")
    return "\n\n".join(parts)


def run_pipeline(uploaded_files, gen_model: str, val_model: str, refine: bool, status):
    client = OpenAI(api_key=OPENAI_API_KEY)
    doc_budget = KNOWN_MODEL_LIMITS[gen_model]["doc_budget"]

    # ── Extract text ──
    status.update(label="Extracting text from documents...")
    docs = []
    for f in uploaded_files:
        text = extract_text(f)
        if text.strip():
            tokens = count_tokens(text)
            docs.append({"name": f.name, "text": text, "tokens": tokens})

    if not docs:
        st.error("No text could be extracted from the uploaded files.")
        return None, []

    total_tokens = sum(d["tokens"] for d in docs)
    st.info(
        f"Loaded **{len(docs)} documents** — "
        f"**{total_tokens:,} tokens** total | "
        f"Model budget: **{doc_budget:,} tokens**"
    )

    # ── Step 1a: Generate initial ontology ──
    if total_tokens <= doc_budget:
        status.update(label=f"Step 1a: Generating ontology ({gen_model}) — direct path...")
        doc_text = build_document_block(docs)
        ttl_string = call_prompt_1a(client, gen_model, doc_text)
    else:
        status.update(label="Documents exceed budget. Chunking...")
        chunks = chunk_all_documents(docs)
        batches = batch_chunks(chunks)
        st.info(f"Split into **{len(chunks)} chunks** across **{len(batches)} batches**.")

        registry: set[str] = set()
        fragments: list[str] = []
        progress = st.progress(0, text="Processing batches...")

        for i, batch in enumerate(batches):
            status.update(label=f"Step 1a: Batch {i + 1}/{len(batches)} ({gen_model})...")
            batch_text = "\n\n---\n\n".join(batch)
            registry_text = ", ".join(sorted(registry)) if registry else ""

            raw_ttl = call_prompt_1a(client, gen_model, batch_text, registry_text)
            is_valid, result = validate_ttl(raw_ttl)

            if is_valid:
                fragments.append(result)
                new_names = extract_registry(result)
                registry.update(new_names)
                logger.info(
                    f"Batch {i + 1}: +{len(new_names)} names, "
                    f"registry total: {len(registry)}"
                )
            else:
                fragments.append(raw_ttl)
                logger.warning(f"Batch {i + 1} produced invalid TTL: {result}")

            progress.progress((i + 1) / len(batches), text=f"Batch {i + 1}/{len(batches)} done — registry: {len(registry)} names")

        status.update(label="Merging ontology fragments...")
        ttl_string = merge_ttl_fragments(fragments)
        st.success(f"Merged {len(fragments)} fragments. Registry: {len(registry)} names.")

    is_valid, result = validate_ttl(ttl_string)
    if is_valid:
        ttl_string = result
        st.success("Step 1a complete: Initial ontology generated.")
    else:
        st.warning("Step 1a: Initial TTL has syntax issues — will attempt to fix.")
        logger.warning(f"Step 1a TTL error: {result}")

    # ── Step 1b: Refine ontology (optional, same model) ──
    if refine:
        status.update(label=f"Step 1b: Refining ontology ({gen_model})...")
        doc_text = build_document_block(docs)
        refined = call_prompt_1b(client, gen_model, doc_text, ttl_string)

        if refined.strip().upper() == "NO_ADDITIONS_NEEDED":
            st.info("Step 1b: Ontology already comprehensive — no changes needed.")
        else:
            refined_valid, refined_result = validate_ttl(refined)
            if refined_valid:
                ttl_string = refined_result
                st.success("Step 1b complete: Ontology refined and corrected.")
            else:
                logger.warning(f"Refinement produced invalid TTL: {refined_result}")
                st.warning("Step 1b: Refinement output had syntax issues. Proceeding with initial ontology.")
    else:
        st.info("Step 1b: Refinement skipped.")

    # ── Step 2: Validate syntax (rdflib first, LLM only if needed) ──
    status.update(label="Step 2: Validating syntax (rdflib)...")
    is_valid, parse_result = validate_ttl(ttl_string)

    if is_valid:
        ttl_string = parse_result
        st.success("Step 2 complete: Syntax validated by rdflib — no LLM call needed.")
    else:
        max_attempts = 5
        current_ttl = ttl_string
        errors = parse_result

        for attempt in range(1, max_attempts + 1):
            st.warning(f"Step 2: Attempt {attempt}/{max_attempts} — calling {val_model} to fix syntax errors...")
            status.update(label=f"Step 2: Fix attempt {attempt}/{max_attempts} ({val_model})...")
            fixed_ttl = call_prompt_2(client, val_model, current_ttl, errors=errors)
            is_fixed, result = validate_ttl(fixed_ttl)

            if is_fixed:
                ttl_string = result
                st.success(f"Step 2 complete: Syntax fixed on attempt {attempt}.")
                break
            else:
                current_ttl = fixed_ttl
                errors = result

            if attempt == max_attempts:
                st.error(f"Step 2: Could not fix syntax errors after {max_attempts} attempts. Last error: {errors}")
                ttl_string = fixed_ttl

    status.update(label="Done!", state="complete")
    return ttl_string, docs


# ---------------------------------------------------------------------------
# Ontology stats
# ---------------------------------------------------------------------------
def get_ontology_stats(ttl_string: str) -> dict:
    g = Graph()
    g.parse(data=ttl_string, format="turtle")
    classes = set(g.subjects(RDF.type, OWL.Class))
    obj_props = set(g.subjects(RDF.type, OWL.ObjectProperty))
    data_props = set(g.subjects(RDF.type, OWL.DatatypeProperty))
    return {
        "Classes": len(classes),
        "Object Properties": len(obj_props),
        "Datatype Properties": len(data_props),
        "Total Triples": len(g),
    }


# ---------------------------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------------------------
def run_kg_pipeline(docs, ttl_string, gen_model, status, alpha=0.1, chunk_size=1000, chunk_overlap=250):
    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
    except Exception as e:
        st.error(f"Could not connect to Neo4j: {e}")
        return

    try:
        # ── Snapshot before build ──
        stats_before = get_graph_stats(driver)

        # ── Convert ontology to schema ──
        status.update(label="Converting ontology to schema...")
        schema = ontology_to_schema(ttl_string)
        st.success(
            f"Schema: **{len(schema.node_types)}** node types, "
            f"**{len(schema.relationship_types)}** relationship types, "
            f"**{len(schema.patterns)}** patterns"
        )

        # ── Build KG ──
        status.update(label=f"Building Knowledge Graph ({gen_model})...")
        progress = st.progress(0, text="Processing documents...")

        def on_doc_complete(done, total, name):
            progress.progress(done / total, text=f"Document {done}/{total}: {name}")

        build_knowledge_graph(driver, schema, docs, gen_model, on_complete=on_doc_complete,
                              chunk_size=chunk_size, chunk_overlap=chunk_overlap)

        # ── Normalize property names (hasName → name safety net) ──
        norm_stats = normalize_entity_names(driver)
        if norm_stats["fixed"]:
            st.info(f"Normalized {norm_stats['fixed']} entity names (hasName → name)")
        st.success("Entity extraction and KG construction complete.")

        # ── Enrich relationships ──
        status.update(label="Enriching relationships (re-scanning chunks)...")
        enrich_progress = st.progress(0, text="Scanning chunks for missing relationships...")

        def on_enrich_progress(done, total):
            enrich_progress.progress(done / total, text=f"Chunk {done}/{total}")

        enrich_stats = enrich_relationships(driver, schema, model=gen_model, on_progress=on_enrich_progress)
        st.success(f"Relationship enrichment: **{enrich_stats['created']}** relationships found across **{enrich_stats['chunks_processed']}** chunks")

        # ── Compute edge weights ──
        status.update(label="Computing edge weights...")
        weight_stats = compute_edge_weights(driver, alpha=alpha)
        st.success(f"Edge weights computed: **{weight_stats['updated']}** relationships (max shared chunks: {weight_stats['max_shared']}, α={alpha})")

        # ── Create evidence layer ──
        status.update(label="Creating contextual evidence layer...")
        ev_stats = create_evidence_layer(driver)
        st.success(f"Evidence layer: **{ev_stats['created']}** evidence nodes created")

        # ── Aggregate confidence ──
        status.update(label="Aggregating confidence scores...")
        agg_stats = aggregate_evidence(driver)
        st.success(f"Confidence aggregated: **{agg_stats['updated']}** relationships scored")

        st.session_state["evidence_ready"] = True

        # ── Snapshot after build ──
        stats_after = get_graph_stats(driver)
        new_entities = stats_after["entities"] - stats_before["entities"]
        new_rels = stats_after["relationships"] - stats_before["relationships"]

        # ── Detect duplicates ──
        status.update(label="Detecting duplicate entities...")
        duplicates = find_duplicate_entities(driver)

        # ── Store results for dashboard ──
        st.session_state["kg_stats"] = {
            "before": stats_before,
            "after": stats_after,
            "new_entities": new_entities,
            "new_relationships": new_rels,
            "duplicates": duplicates,
        }

        status.update(label="Done!", state="complete")

    finally:
        neo4j_client.close()


def main():
    st.set_page_config(page_title="Ontology-Driven KG Builder", layout="wide")
    st.title("Ontology-Driven Knowledge Graph Builder")
    st.caption("Upload documents → Generate ontology → Build Knowledge Graph")

    # ── Sidebar ──
    model_names = list(KNOWN_MODEL_LIMITS.keys())
    with st.sidebar:
        st.header("Configuration")
        gen_model = st.selectbox("Generation Model", options=model_names, index=0)
        val_model = st.selectbox("Syntax Fix Model", options=model_names, index=1,
                                 help="Only used if rdflib detects syntax errors")
        st.divider()
        refine = st.checkbox("Refine ontology (Step 1b)", value=False,
                             help="Extra LLM pass to add missing entity types and relationships")
        budget = KNOWN_MODEL_LIMITS[gen_model]["doc_budget"]
        st.metric("Document Budget", f"{budget:,} tokens")
        st.divider()
        alpha = st.slider("Edge Weight Floor (α)", min_value=0.0, max_value=0.5, value=0.1, step=0.05,
                          help="Minimum weight for any relationship. Higher = less contrast between weak and strong edges.")
        st.divider()
        st.header("Extraction")
        chunk_size = st.slider("Chunk Size (tokens)", min_value=500, max_value=3000, value=1000, step=100,
                               help="Smaller chunks = more focused extraction = more relationships found")
        chunk_overlap = st.slider("Chunk Overlap (tokens)", min_value=50, max_value=500, value=250, step=50,
                                  help="Overlap between chunks to preserve context at boundaries")
        st.divider()
        st.header("Knowledge Graph"
        )

    # ── File uploader ──
    uploaded_files = st.file_uploader(
        "Upload Documents",
        type=ACCEPTED_FILE_TYPES,
        accept_multiple_files=True,
        help="Supported formats: .txt, .pdf, .docx — upload related documents only",
    )

    if uploaded_files:
        file_info = []
        for f in uploaded_files:
            content = f.read()
            f.seek(0)
            size_kb = len(content) / 1024
            file_info.append({"File": f.name, "Size": f"{size_kb:.1f} KB"})
        st.dataframe(file_info, use_container_width=True, hide_index=True)

    # ── Step 1: Generate Ontology ──
    st.header("Step 1: Generate Ontology")

    if st.button("Generate Ontology", type="primary", disabled=not uploaded_files):
        with st.status("Starting ontology pipeline...", expanded=True) as status:
            ttl_result, docs = run_pipeline(uploaded_files, gen_model, val_model, refine, status)

        if ttl_result:
            st.session_state["ttl_result"] = ttl_result
            st.session_state["docs"] = docs

    ttl_result = st.session_state.get("ttl_result")
    docs = st.session_state.get("docs")

    if ttl_result:
        st.subheader("Ontology Summary")
        try:
            stats = get_ontology_stats(ttl_result)
            cols = st.columns(len(stats))
            for col, (label, value) in zip(cols, stats.items()):
                col.metric(label, value)
        except Exception:
            pass

        with st.expander("View Generated Ontology (.ttl)"):
            st.code(ttl_result, language="turtle")

        st.download_button(
            label="Download ontology.ttl",
            data=ttl_result,
            file_name="ontology.ttl",
            mime="text/turtle",
        )

        # ── Step 2: Build Knowledge Graph ──
        st.header("Step 2: Build Knowledge Graph")

        if st.button("Build Knowledge Graph", type="primary"):
            with st.status("Building Knowledge Graph...", expanded=True) as kg_status:
                run_kg_pipeline(docs, ttl_result, gen_model, kg_status, alpha=alpha,
                                chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    # ── Overlap Dashboard ──
    kg_stats = st.session_state.get("kg_stats")
    if kg_stats:
        st.divider()
        st.header("Overlap Dashboard")

        col1, col2, col3, col4 = st.columns(4)
        col1.metric("Entities Before", kg_stats["before"]["entities"])
        col2.metric("New Entities", kg_stats["new_entities"])
        col3.metric("Total Entities", kg_stats["after"]["entities"])
        col4.metric("Duplicate Pairs", len(kg_stats["duplicates"]))

        duplicates = kg_stats["duplicates"]
        if duplicates:
            same_label = [d for d in duplicates if d["match_type"] == "same_label"]
            cross_label = [d for d in duplicates if d["match_type"] == "cross_label"]
            overlap_pct = (len(duplicates) / kg_stats["after"]["entities"]) * 100 if kg_stats["after"]["entities"] else 0
            st.warning(f"**{len(duplicates)} duplicate entity pairs found** ({overlap_pct:.1f}% of total entities)")

            if same_label:
                with st.expander(f"Same-type duplicates — high confidence ({len(same_label)})"):
                    dup_display = [{"Name": d["name"], "Type": d["label_a"]} for d in same_label]
                    st.dataframe(dup_display, use_container_width=True, hide_index=True)

            if cross_label:
                with st.expander(f"Cross-type duplicates — needs review ({len(cross_label)})"):
                    dup_display = [{"Name": d["name"], "Type A": d["label_a"], "Type B": d["label_b"]} for d in cross_label]
                    st.dataframe(dup_display, use_container_width=True, hide_index=True)
        else:
            st.success("No duplicate entities found. Graph is clean.")

    # ── Entity Resolution ──
    run_entity_resolution_section()

    # ── Contextual Evidence ──
    run_contextual_enrichment_section(gen_model)

    # ── External Web Sources ──
    run_web_sources_section(gen_model)

    # ── Floating Chatbot ──
    run_floating_chatbot(gen_model)


CONFIDENCE_COLORS = {"high": "#2ecc71", "medium": "#f39c12", "low": "#e74c3c", "error": "#95a5a6"}


def run_entity_resolution_section():
    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            neo4j_client.close()
            return
    except Exception:
        return

    st.divider()
    st.header("Entity Resolution")

    # ── Scoring controls ──
    st.subheader("1. Find & Score Candidates")
    score_col1, score_col2 = st.columns(2)
    with score_col1:
        scoring_level = st.radio(
            "Scoring level",
            ["Exact match", "Exact + Embedding", "Exact + Embedding + LLM"],
            index=0,
            help="Higher levels catch more fuzzy duplicates",
        )
    with score_col2:
        llm_model = st.selectbox("LLM Judge Model", ["gpt-4o-mini", "gpt-4o"], index=0,
                                  help="Only used at LLM scoring level")

    if st.button("Run Entity Resolution Scoring", type="primary"):
        with st.spinner("Finding candidate pairs..."):
            candidates = find_candidate_pairs(driver)
            all_pairs = candidates["same_label"] + candidates["cross_label"]

        if not all_pairs:
            st.success("No duplicate candidates found. Graph is clean.")
            neo4j_client.close()
            return

        st.info(f"Found **{len(candidates['same_label'])}** same-type and **{len(candidates['cross_label'])}** cross-type candidate pairs.")

        # Level 1: Exact
        scored = [score_exact(p) for p in all_pairs]

        # Level 2: Embedding
        if scoring_level in ("Exact + Embedding", "Exact + Embedding + LLM"):
            with st.spinner("Computing embedding similarity..."):
                scored = score_embedding_batch(all_pairs)

        # Level 3: LLM judge (only for ambiguous pairs)
        if scoring_level == "Exact + Embedding + LLM":
            ambiguous = [s for s in scored if s["confidence"] in ("medium", "low")]
            if ambiguous:
                with st.spinner(f"LLM judging {len(ambiguous)} ambiguous pairs..."):
                    llm_scored = score_llm_batch(ambiguous, model=llm_model)
                    llm_map = {(s["id_a"], s["id_b"]): s for s in llm_scored}
                    scored = [llm_map.get((s["id_a"], s["id_b"]), s) for s in scored]

        st.session_state["er_scored_pairs"] = scored

    # ── Per-pair approval ──
    scored_pairs = st.session_state.get("er_scored_pairs")
    if scored_pairs:
        st.subheader("2. Review & Approve Pairs")
        st.caption("Check the pairs you want to merge. Uncheck to keep separate.")

        sorted_pairs = sorted(scored_pairs, key=lambda x: (
            {"high": 0, "medium": 1, "low": 2, "error": 3}.get(x["confidence"], 3),
            -x["score"],
        ))

        if "er_approvals" not in st.session_state:
            st.session_state["er_approvals"] = {
                i: p["confidence"] == "high" for i, p in enumerate(sorted_pairs)
            }

        for i, pair in enumerate(sorted_pairs):
            color = CONFIDENCE_COLORS.get(pair["confidence"], "#aaa")
            col_check, col_info, col_detail = st.columns([0.5, 3, 4])

            with col_check:
                st.session_state["er_approvals"][i] = st.checkbox(
                    "Merge", value=st.session_state["er_approvals"].get(i, False),
                    key=f"er_pair_{i}", label_visibility="collapsed",
                )

            with col_info:
                st.markdown(
                    f'<span style="color:{color}; font-weight:bold;">{pair["confidence"].upper()}</span> '
                    f'({pair["score_type"]}: {pair["score"]:.2f})  **{pair["name"]}**',
                    unsafe_allow_html=True,
                )

            with col_detail:
                type_info = (f'{pair["label_a"]}' if pair["label_a"] == pair["label_b"]
                             else f'{pair["label_a"]} vs {pair["label_b"]}')
                st.caption(f'{type_info} — {pair["reason"]}')

        approved = [sorted_pairs[i] for i, checked in st.session_state["er_approvals"].items() if checked]
        st.info(f"**{len(approved)}** of **{len(sorted_pairs)}** pairs approved for merge.")

        # ── Clustering + Merge ──
        st.subheader("3. Merge")

        if not approved:
            st.warning("No pairs approved. Check at least one pair above to enable merge.")
        else:
            clusters = build_transitive_clusters(approved)
            oversized = [c for c in clusters if len(c) > 5]

            with st.expander(f"Preview: {len(clusters)} merge clusters"):
                for ci, cluster in enumerate(clusters):
                    names = [f'{n["name"]} ({n["label"]})' for n in cluster]
                    survivor = names[0]
                    icon = "!" if len(cluster) > 5 else str(ci + 1)
                    st.markdown(f"**Cluster {icon}:** Keep **{survivor}**, merge: {', '.join(names[1:])}")

            if oversized:
                st.warning(f"{len(oversized)} cluster(s) exceed 5 members — review carefully before merging.")

            if st.button("Execute Merge", type="primary"):
                merge_log = []
                progress = st.progress(0, text="Merging clusters...")

                for ci, cluster in enumerate(clusters):
                    result = merge_cluster(driver, cluster)
                    merge_log.append(result)
                    progress.progress((ci + 1) / len(clusters), text=f"Merged cluster {ci + 1}/{len(clusters)}")

                st.session_state["merge_log"] = merge_log

                total_merged = sum(r["dropped_count"] for r in merge_log if r["status"] == "merged")
                st.success(f"Merged {total_merged} duplicate nodes across {len(clusters)} clusters.")

                # Refresh stats
                stats_after = get_graph_stats(driver)
                st.session_state["kg_stats"]["after"] = stats_after
                st.session_state["kg_stats"]["duplicates"] = find_duplicate_entities(driver)

                # Check for orphans
                orphans = find_orphaned_nodes(driver)
                if orphans:
                    st.warning(f"{len(orphans)} orphaned node(s) detected after merge:")
                    st.dataframe(
                        [{"Name": o["name"], "Type": o["label"]} for o in orphans],
                        use_container_width=True, hide_index=True,
                    )

                # Clear scored pairs since graph changed
                st.session_state.pop("er_scored_pairs", None)
                st.session_state.pop("er_approvals", None)

    # ── Undo ──
    merge_log = st.session_state.get("merge_log")
    if merge_log:
        all_snapshots = []
        for entry in merge_log:
            if entry.get("snapshots"):
                all_snapshots.extend(entry["snapshots"])

        if all_snapshots:
            st.divider()
            if st.button("Undo Last Merge", type="secondary"):
                with st.spinner("Restoring merged nodes..."):
                    restored = undo_merge(driver, all_snapshots)
                    st.success(f"Restored {restored} node(s). Note: re-run scoring to verify graph state.")
                    st.session_state.pop("merge_log", None)

    neo4j_client.close()


def run_contextual_enrichment_section(gen_model: str):
    st.divider()
    st.header("Contextual Evidence Layer")
    st.caption("Enterprise-style provenance, confidence, and temporal metadata on every relationship")

    if not st.session_state.get("evidence_ready"):
        st.info("Build a Knowledge Graph first — the evidence layer is created automatically.")
        return

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No entities in graph. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    # ── Evidence dashboard ──
    try:
        stats = get_evidence_stats(driver)
        col1, col2, col3, col4 = st.columns(4)
        col1.metric("Evidence Nodes", stats["total_evidence"])
        col2.metric("Avg Confidence", f"{stats['avg_confidence']:.3f}")
        col3.metric("Temporal Enriched", stats["temporal_enriched"])
        col4.metric("Relationships Scored", stats["relationships_scored"])

        if stats["by_source_type"]:
            with st.expander("Evidence by source type"):
                for src_type, cnt in stats["by_source_type"].items():
                    st.markdown(f"- **{src_type}**: {cnt} nodes")
    except Exception as e:
        logger.warning(f"Could not load evidence stats: {e}")

    # ── Temporal enrichment (LLM cost) ──
    st.subheader("Temporal Enrichment")
    st.caption("Use LLM to determine when each relationship was valid (costs API tokens)")

    tcol1, tcol2 = st.columns(2)
    with tcol1:
        temporal_model = st.selectbox("Temporal Model", ["gpt-4o-mini", "gpt-4o"], index=0,
                                       key="temporal_model")
    with tcol2:
        temporal_batch = st.slider("Batch size", 5, 30, 15, key="temporal_batch",
                                    help="Relationships per LLM call")

    if st.button("Run Temporal Enrichment", type="primary", key="btn_temporal"):
        with st.status("Enriching temporal metadata...", expanded=True) as t_status:
            progress = st.progress(0, text="Processing batches...")

            def on_temporal_progress(done, total):
                progress.progress(done / total, text=f"Batch {done}/{total}")

            result = enrich_temporal(driver, model=temporal_model,
                                     batch_size=temporal_batch,
                                     on_progress=on_temporal_progress)
            st.success(f"Temporal enrichment: **{result['updated']}** evidence nodes updated across **{result['batches']}** batches")

            t_status.update(label="Re-aggregating confidence with temporal data...")
            aggregate_evidence(driver)
            t_status.update(label="Done!", state="complete")

    # ── Web evidence (after web sources are fetched) ──
    if st.session_state.get("web_sources_enabled"):
        st.subheader("Web Evidence")
        if st.button("Add Web Evidence", key="btn_web_evidence"):
            with st.spinner("Creating evidence from web sources..."):
                web_ev = create_web_evidence(driver)
                st.success(f"Web evidence: **{web_ev['created']}** evidence nodes created")
                aggregate_evidence(driver)
                st.success("Confidence re-aggregated with web evidence.")

    # ── Maintenance ──
    with st.expander("Maintenance"):
        mcol1, mcol2 = st.columns(2)
        with mcol1:
            if st.button("Re-aggregate Confidence", key="btn_reagg"):
                with st.spinner("Aggregating..."):
                    result = aggregate_evidence(driver)
                    st.success(f"Re-aggregated: {result['updated']} relationships")
        with mcol2:
            if st.button("Clear Evidence Layer", type="secondary", key="btn_clear_ev"):
                with st.spinner("Clearing..."):
                    result = clear_evidence_layer(driver)
                    st.success(f"Cleared: {result['deleted']} evidence nodes removed")
                    st.session_state["evidence_ready"] = False

    neo4j_client.close()


def run_web_sources_section(gen_model: str):
    st.divider()
    st.header("External Web Sources")
    st.caption("Augment your knowledge graph with related web content")

    ttl_result = st.session_state.get("ttl_result")
    docs = st.session_state.get("docs")
    if not ttl_result or not docs:
        st.info("Generate an ontology and build a Knowledge Graph first.")
        return

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No entities in graph. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    # Show existing web stats if present
    try:
        ws = get_web_source_stats(driver)
        if ws["web_chunks"] > 0:
            st.session_state["web_sources_enabled"] = True
            st.success(f"Web content active: **{ws['web_chunks']}** web chunks, **{ws['similar_to_edges']}** SIMILAR_TO edges (avg weight: {ws['avg_edge_weight']:.3f})")
    except Exception:
        pass

    col1, col2, col3 = st.columns(3)
    with col1:
        max_articles = st.slider("Max articles", 1, 10, 5, key="ws_max_articles")
    with col2:
        sim_threshold = st.slider("Similarity threshold", 0.3, 0.8, 0.5, 0.05,
                                  key="ws_sim_threshold",
                                  help="Min cosine similarity to create SIMILAR_TO edge")
    with col3:
        max_topics = st.slider("Max topics", 2, 7, 5, key="ws_max_topics")

    if st.button("Fetch & Process Web Sources", type="primary"):
        with st.status("Processing web sources...", expanded=True) as ws_status:
            ws_status.update(label="Extracting key topics from your documents...")
            topics = extract_topics(docs, model=gen_model, max_topics=max_topics)
            st.info(f"Topics: {', '.join(topics)}")

            ws_status.update(label="Searching the web via OpenAI...")
            web_docs = search_and_fetch(topics[:max_articles], model=gen_model)

            if not web_docs:
                st.warning("No content found. Try different documents or topics.")
                ws_status.update(label="No results", state="error")
                neo4j_client.close()
                return

            st.info(f"Fetched web content for **{len(web_docs)}** topics")
            with st.expander("Web search results"):
                for r in web_docs:
                    st.markdown(f"- **{r['name']}** — {len(r['text'])} chars")

            ws_status.update(label=f"Building web knowledge graph ({gen_model})...")
            schema = ontology_to_schema(ttl_result)
            progress = st.progress(0, text="Processing web articles...")

            def on_web_complete(done, total, name):
                progress.progress(done / total, text=f"Web article {done}/{total}: {name}")

            build_web_knowledge_graph(driver, schema, web_docs, gen_model, on_complete=on_web_complete)

            ws_status.update(label="Computing SIMILAR_TO edges...")
            edge_stats = compute_similar_to_edges(driver, similarity_threshold=sim_threshold)

            st.session_state["web_sources_enabled"] = True
            st.session_state["web_source_stats"] = edge_stats

            ws_status.update(label="Done!", state="complete")

    ws_stats = st.session_state.get("web_source_stats")
    if ws_stats:
        mcol1, mcol2, mcol3 = st.columns(3)
        mcol1.metric("SIMILAR_TO Edges", ws_stats["edges_created"])
        mcol2.metric("Avg Similarity", f"{ws_stats['avg_similarity']:.3f}")
        mcol3.metric("Max Similarity", f"{ws_stats['max_similarity']:.3f}")

    if st.session_state.get("web_sources_enabled"):
        if st.button("Remove All Web Content", type="secondary"):
            with st.spinner("Removing web content..."):
                result = remove_web_content(driver)
                st.success(
                    f"Removed {result['web_chunks_removed']} web chunks, "
                    f"{result['web_documents_removed']} web documents, "
                    f"{result['edges_removed']} SIMILAR_TO edges."
                )
                st.session_state["web_sources_enabled"] = False
                st.session_state.pop("web_source_stats", None)

    neo4j_client.close()


def run_floating_chatbot(gen_model: str):
    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = []

    with st.sidebar:
        st.divider()
        st.header("GraphRAG Chat")
        web_available = st.session_state.get("web_sources_enabled", False)
        include_web = st.checkbox("Include web sources", value=False, key="chat_web",
                                  disabled=not web_available)
        if st.button("Clear chat", key="btn_clear_chat"):
            st.session_state["chat_history"] = []
            st.rerun()

    for msg in st.session_state["chat_history"]:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
            if msg.get("context_items"):
                with st.expander("Retrieved context"):
                    for ci in msg["context_items"]:
                        st.markdown(f"**Chunk {ci['index']}** (score: {ci['score']})")
                        st.text(ci["text"])
                        st.divider()

    question = st.chat_input("Ask a question about your knowledge graph...")

    if question:
        st.session_state["chat_history"].append({"role": "user", "content": question})

        try:
            neo4j_client = Neo4jClient()
            driver = neo4j_client()

            if not has_any_entities(driver):
                answer = "No knowledge graph found. Please build one first."
                st.session_state["chat_history"].append({"role": "assistant", "content": answer})
                neo4j_client.close()
                st.rerun()
                return

            with st.spinner("Searching knowledge graph..."):
                result = query_graph_rag(
                    driver, question, gen_model,
                    hops=2,
                    weight_threshold=0.1,
                    confidence_threshold=0.0,
                    include_web_sources=include_web,
                )
            neo4j_client.close()

            answer = result["answer"]
            context_items = []
            if result["context"] and result["context"].items:
                for i, item in enumerate(result["context"].items, 1):
                    content_str = str(item.content)
                    score_match = re.search(r'score=([\d.]+)', content_str)
                    score_val = f"{float(score_match.group(1)):.4f}" if score_match else "N/A"
                    context_items.append({"index": i, "score": score_val, "text": content_str[:500]})

            st.session_state["chat_history"].append({
                "role": "assistant",
                "content": answer,
                "context_items": context_items,
            })
            st.rerun()

        except Exception as e:
            error_msg = f"Error querying graph: {e}"
            st.session_state["chat_history"].append({"role": "assistant", "content": error_msg})
            st.rerun()


if __name__ == "__main__":
    main()
